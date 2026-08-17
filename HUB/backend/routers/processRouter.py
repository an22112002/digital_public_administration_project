from fastapi import APIRouter, UploadFile, File, HTTPException
from backend.services.processServices import getActiveServices
from openpyxl import load_workbook
from io import BytesIO

process_router = APIRouter(prefix="/process", tags=["process"])

@process_router.post("/services/import")
async def import_xlsx(file: UploadFile = File(...)):
    if not file.filename.endswith(".xlsx"):
        raise HTTPException(
            status_code=400,
            detail="Chỉ file .xlsx được chấp nhận"
        )

    file_data = await file.read()

    try:
        workbook = load_workbook(
            filename=BytesIO(file_data),
            read_only=True,
            data_only=True
        )
    except Exception:
        raise HTTPException(
            status_code=400,
            detail="Invalid XLSX file"
        )

    print(workbook.sheetnames)

    for sheet_name in workbook.sheetnames:
        worksheet = workbook[sheet_name]

        rows = worksheet.iter_rows(values_only=True)

        # dòng đầu tiên = header
        headers = next(rows, None)

        if not headers:
            continue

        headers = [
            str(header).strip()
            if header is not None
            else None
            for header in headers
        ]

        for row in rows:
            data = dict(zip(headers, row))

            print(sheet_name, data)

    return {
        "success": True,
        "sheets": workbook.sheetnames
    }
    
@process_router.get("/services/active")
def getActiveServices():
    result = getActiveServices()
    return result