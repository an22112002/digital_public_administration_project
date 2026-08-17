from io import BytesIO

from fastapi import HTTPException
from openpyxl import load_workbook

from database.index import db

from backend.models.serviceModels import (
    ServiceImport,
    DefaultDocumentImport,
    ExtentionDocumentImport,
)

from backend.crud.servicesCrud import (
    insert_services,
    get_service_ids,
    delete_all_services,
)

from backend.crud.defaultDocumentCrud import (
    insert_default_documents,
    delete_all_default_documents,
)

from backend.crud.extentionDocumentCrud import (
    insert_extention_documents,
    delete_all_extention_documents,
)


# Tên các sheet bắt buộc phải có trong file Excel.
REQUIRED_SHEETS = {
    "services",
    "default",
    "extention",
}


def import_xlsx(file_data: bytes) -> dict:
    """
    Import toàn bộ dữ liệu từ file XLSX vào database.

    File Excel gồm 3 sheet:

    - services
    - default
    - extention

    Toàn bộ quá trình import sử dụng cùng một cursor
    để đảm bảo các thao tác database nằm trong cùng một transaction.
    """

    workbook = _load_workbook(file_data)

    _validate_sheets(workbook)

    # Đọc dữ liệu từ từng sheet.
    services = _parse_services(workbook["services"])

    # Mở connection/cursor để thao tác database.
    with db.get_cursor() as cursor:

        # Xóa dữ liệu cũ trước khi import.
        # Vì đây là dữ liệu cấu hình nên Excel được xem là source of truth.
        _clear_old_data(cursor)

        # Insert services trước.
        insert_services(
            cursor,
            services
        )

        # Lấy serviceID vừa được tạo.
        service_ids = get_service_ids(cursor)

        # Parse default documents.
        default_documents = _parse_default_documents(
            workbook["default"],
            service_ids
        )

        # Insert default documents.
        insert_default_documents(
            cursor,
            default_documents
        )

        # Parse extention documents.
        extention_documents = _parse_extention_documents(
            workbook["extention"]
        )

        # Insert extention documents.
        insert_extention_documents(
            cursor,
            extention_documents
        )

    return {
        "services": len(services),
        "default_documents": len(default_documents),
        "extention_documents": len(extention_documents),
    }


def _load_workbook(file_data: bytes):
    """
    Đọc bytes của file XLSX và trả về workbook.
    """

    try:
        return load_workbook(
            filename=BytesIO(file_data),
            read_only=True,
            data_only=True
        )
    except Exception as exc:
        raise HTTPException(
            status_code=400,
            detail=f"File Excel không hợp lệ: {exc}"
        )


def _validate_sheets(workbook) -> None:
    """
    Kiểm tra file Excel có đầy đủ các sheet bắt buộc hay không.
    """

    actual_sheets = set(workbook.sheetnames)

    missing_sheets = REQUIRED_SHEETS - actual_sheets

    if missing_sheets:
        raise HTTPException(
            status_code=400,
            detail=(
                "File Excel thiếu sheet: "
                + ", ".join(sorted(missing_sheets))
            )
        )


def _get_headers(worksheet) -> list[str]:
    """
    Đọc dòng đầu tiên của worksheet làm tên các trường.
    """

    rows = worksheet.iter_rows(values_only=True)

    headers = next(rows, None)

    if not headers:
        raise HTTPException(
            status_code=400,
            detail=f"Sheet '{worksheet.title}' không có header"
        )

    return [
        str(header).strip()
        if header is not None
        else ""
        for header in headers
    ]


def _validate_headers(
    worksheet,
    headers: list[str],
    required_headers: set[str]
) -> None:
    """
    Kiểm tra worksheet có đầy đủ các cột bắt buộc.
    """

    actual_headers = set(headers)

    missing_headers = required_headers - actual_headers

    if missing_headers:
        raise HTTPException(
            status_code=400,
            detail=(
                f"Sheet '{worksheet.title}' thiếu cột: "
                + ", ".join(sorted(missing_headers))
            )
        )


def _parse_services(worksheet) -> list[ServiceImport]:
    """
    Đọc và convert sheet services thành danh sách ServiceImport.
    """

    headers = _get_headers(worksheet)

    _validate_headers(
        worksheet,
        headers,
        {
            "title",
            "realTitle",
            "url",
            "buttonPosition",
            "processes",
            "active",
        }
    )

    rows = worksheet.iter_rows(values_only=True)

    # Bỏ qua dòng header.
    next(rows)

    result = []

    for row_number, row in enumerate(rows, start=2):

        # Bỏ qua dòng hoàn toàn trống.
        if _is_empty_row(row):
            continue

        data = dict(zip(headers, row))

        title = data.get("title")
        real_title = data.get("realTitle")
        url = data.get("url")

        if not title:
            _raise_row_error(
                worksheet,
                row_number,
                "thiếu title"
            )

        if not real_title:
            _raise_row_error(
                worksheet,
                row_number,
                "thiếu realTitle"
            )

        if not url:
            _raise_row_error(
                worksheet,
                row_number,
                "thiếu url"
            )

        result.append(
            ServiceImport(
                title=str(title).strip(),
                realTitle=str(real_title).strip(),
                url=str(url).strip(),

                buttonPosition=_parse_int(
                    data.get("buttonPosition"),
                    default=1
                ),

                processes=_parse_processes(
                    data.get("processes")
                ),

                active=_parse_bool(
                    data.get("active"),
                    default=True
                )
            )
        )

    return result


def _parse_default_documents(
    worksheet,
    service_ids: dict[str, int]
) -> list[DefaultDocumentImport]:
    """
    Đọc sheet default và chuyển realServiceTitle
    thành serviceID tương ứng trong database.
    """

    headers = _get_headers(worksheet)

    _validate_headers(
        worksheet,
        headers,
        {
            "realServiceTitle",
            "name",
            "description",
            "required",
            "OCRtab",
        }
    )

    rows = worksheet.iter_rows(values_only=True)

    # Bỏ qua dòng header.
    next(rows)

    result = []

    for row_number, row in enumerate(rows, start=2):

        if _is_empty_row(row):
            continue

        data = dict(zip(headers, row))

        real_service_title = data.get(
            "realServiceTitle"
        )

        if not real_service_title:
            _raise_row_error(
                worksheet,
                row_number,
                "thiếu realServiceTitle"
            )

        real_service_title = str(
            real_service_title
        ).strip()

        service_id = service_ids.get(
            real_service_title
        )

        if service_id is None:
            _raise_row_error(
                worksheet,
                row_number,
                (
                    "không tìm thấy service có "
                    f"realTitle='{real_service_title}'"
                )
            )

        name = data.get("name")

        if not name:
            _raise_row_error(
                worksheet,
                row_number,
                "thiếu name"
            )

        result.append(
            DefaultDocumentImport(
                serviceID=service_id,
                name=str(name).strip(),

                description=_parse_string(
                    data.get("description")
                ),

                required=_parse_bool(
                    data.get("required"),
                    default=False
                ),

                OCRtab=_parse_string(
                    data.get("OCRtab")
                )
            )
        )

    return result


def _parse_extention_documents(
    worksheet
) -> list[ExtentionDocumentImport]:
    """
    Đọc sheet extention và chuyển thành danh sách
    ExtentionDocumentImport.
    """

    headers = _get_headers(worksheet)

    _validate_headers(
        worksheet,
        headers,
        {
            "name",
            "description",
            "OCRtab",
        }
    )

    rows = worksheet.iter_rows(values_only=True)

    # Bỏ qua dòng header.
    next(rows)

    result = []

    for row_number, row in enumerate(rows, start=2):

        if _is_empty_row(row):
            continue

        data = dict(zip(headers, row))

        name = data.get("name")

        if not name:
            _raise_row_error(
                worksheet,
                row_number,
                "thiếu name"
            )

        result.append(
            ExtentionDocumentImport(
                name=str(name).strip(),

                description=_parse_string(
                    data.get("description")
                ),

                OCRtab=_parse_string(
                    data.get("OCRtab")
                )
            )
        )

    return result


def _parse_processes(value) -> list[str]:
    """
    Chuyển giá trị processes trong Excel thành list[str].

    Ví dụ:
    'step1,step2,step3'

    sẽ thành:

    ['step1', 'step2', 'step3']
    """

    if value is None:
        return []

    if isinstance(value, list):
        return [
            str(item).strip()
            for item in value
            if item
        ]

    return [
        item.strip()
        for item in str(value).split(",")
        if item.strip()
    ]


def _parse_bool(
    value,
    default: bool
) -> bool:
    """
    Chuyển giá trị Excel thành boolean.

    Các giá trị được coi là True:
    - true
    - 1
    - yes
    - y

    Nếu giá trị rỗng thì trả về default.
    """

    if value is None:
        return default

    if isinstance(value, bool):
        return value

    if isinstance(value, (int, float)):
        return value != 0

    return str(value).strip().lower() in {
        "true",
        "1",
        "yes",
        "y",
    }


def _parse_int(
    value,
    default: int
) -> int:
    """
    Chuyển giá trị Excel thành integer.
    """

    if value is None:
        return default

    try:
        return int(value)
    except (ValueError, TypeError):
        return default


def _parse_string(value) -> str | None:
    """
    Chuyển giá trị Excel thành string.

    Giá trị None hoặc chuỗi rỗng sẽ trả về None.
    """

    if value is None:
        return None

    value = str(value).strip()

    return value or None


def _is_empty_row(row) -> bool:
    """
    Kiểm tra một dòng Excel có hoàn toàn trống hay không.
    """

    return all(
        value is None
        for value in row
    )


def _raise_row_error(
    worksheet,
    row_number: int,
    message: str
) -> None:
    """
    Tạo HTTPException chứa thông tin sheet và dòng bị lỗi.
    """

    raise HTTPException(
        status_code=400,
        detail=(
            f"Sheet '{worksheet.title}', "
            f"dòng {row_number}: {message}"
        )
    )


def _clear_old_data(cursor) -> None:
    """
    Xóa dữ liệu cũ trước khi import.

    Xóa theo thứ tự:
    1. default_documents
    2. services
    3. extention_documents

    default_documents có foreign key tới services
    nên phải xóa default_documents trước services.
    """

    delete_all_default_documents(cursor)

    delete_all_services(cursor)

    delete_all_extention_documents(cursor)